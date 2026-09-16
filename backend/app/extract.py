"""Turn an uploaded complaint document into plain text.

Deliberately simple: the assignment states production-grade OCR is not required.
PDFs with a text layer, DOCX, EML and TXT cover the realistic complaint formats.
"""
import email
import io
from email import policy

MAX_CHARS = 20_000


def _pdf(raw: bytes) -> str:
    import pymupdf

    with pymupdf.open(stream=raw, filetype="pdf") as doc:
        return "\n".join(page.get_text() for page in doc)


def _docx(raw: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(raw))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _eml(raw: bytes) -> str:
    msg = email.message_from_bytes(raw, policy=policy.default)
    body = msg.get_body(preferencelist=("plain", "html"))
    text = body.get_content() if body else ""
    return f"From: {msg.get('From','')}\nSubject: {msg.get('Subject','')}\nDate: {msg.get('Date','')}\n\n{text}"


def document_to_text(filename: str, raw: bytes) -> str:
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        text = _pdf(raw)
    elif name.endswith(".docx"):
        text = _docx(raw)
    elif name.endswith(".eml"):
        text = _eml(raw)
    else:
        text = raw.decode("utf-8", errors="replace")

    text = text.strip()
    if not text:
        raise ValueError(
            "No readable text found in this file. Scanned images need OCR, "
            "which is out of scope for this module."
        )
    return text[:MAX_CHARS]
